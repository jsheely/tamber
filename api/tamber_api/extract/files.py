"""Document extractors (docs/API.md §8.5 c): PDF, DOCX, EPUB, HTML, Markdown, plain text.

Type detection: declared Content-Type first, then the file extension, then magic bytes.
pypdf (BSD) is used for PDF; pymupdf (AGPL) is deliberately not used.
"""

from __future__ import annotations

import io
import logging
import re
import zipfile
from dataclasses import dataclass
from pathlib import PurePosixPath
from typing import Any

from bs4 import BeautifulSoup
from charset_normalizer import from_bytes

from tamber_api.errors import TamberError
from tamber_api.extract.html import block_paragraphs, extract_html, first_heading
from tamber_api.extract.normalize import clean_paragraph, split_plain_paragraphs

logger = logging.getLogger("tamber.extract")

PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
EPUB = "application/epub+zip"
HTML = "text/html"
MARKDOWN = "text/markdown"
TEXT = "text/plain"

_CONTENT_TYPES = {
    "application/pdf": PDF,
    "application/x-pdf": PDF,
    DOCX: DOCX,
    "application/epub+zip": EPUB,
    "text/html": HTML,
    "application/xhtml+xml": HTML,
    "text/markdown": MARKDOWN,
    "text/x-markdown": MARKDOWN,
    "text/plain": TEXT,
}
_EXTENSIONS = {
    ".pdf": PDF,
    ".docx": DOCX,
    ".epub": EPUB,
    ".html": HTML,
    ".htm": HTML,
    ".xhtml": HTML,
    ".md": MARKDOWN,
    ".markdown": MARKDOWN,
    ".txt": TEXT,
    ".text": TEXT,
}


@dataclass(slots=True)
class Extracted:
    title: str | None
    paragraphs: list[str]
    language: str | None
    mime_type: str


def _zip_kind(data: bytes) -> str | None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            names = set(zf.namelist())
            if "mimetype" in names:
                mimetype = zf.read("mimetype").decode("ascii", "ignore").strip()
                if mimetype == EPUB:
                    return EPUB
            if "word/document.xml" in names:
                return DOCX
    except (zipfile.BadZipFile, OSError, KeyError):
        return None
    return None


def _sniff(data: bytes) -> str | None:
    head = data[:1024].lstrip()
    if data[:5] == b"%PDF-":
        return PDF
    if data[:4] == b"PK\x03\x04":
        return _zip_kind(data)
    lowered = head[:256].lower()
    if lowered.startswith((b"<!doctype html", b"<html")) or b"<html" in lowered[:128]:
        return HTML
    if b"\x00" in data[:4096]:
        return None
    return TEXT


def detect_type(data: bytes, content_type: str | None, filename: str | None) -> str | None:
    """Declared type, then extension, then magic bytes. None = unsupported."""
    declared = (content_type or "").split(";")[0].strip().lower()
    if declared in _CONTENT_TYPES:
        kind = _CONTENT_TYPES[declared]
        # A zip claiming to be a document must actually be one.
        if kind in (DOCX, EPUB) and _zip_kind(data) != kind:
            return _sniff(data)
        return kind
    if filename:
        ext = PurePosixPath(filename.replace("\\", "/")).suffix.lower()
        if ext in _EXTENSIONS:
            kind = _EXTENSIONS[ext]
            if kind in (DOCX, EPUB) and _zip_kind(data) != kind:
                return _sniff(data)
            return kind
    if declared and declared not in ("application/octet-stream", "binary/octet-stream", ""):
        # A declared type we don't read; let magic bytes decide for PDFs/zips served oddly.
        sniffed = _sniff(data)
        return sniffed if sniffed in (PDF, DOCX, EPUB) else None
    return _sniff(data)


def decode_text(data: bytes) -> str:
    if data.startswith(b"\xef\xbb\xbf"):
        data = data[3:]
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        pass
    best = from_bytes(data).best()
    if best is None:
        return data.decode("utf-8", errors="replace")
    return str(best)


# --- per-format extractors -----------------------------------------------------------------------

_DEHYPHEN_RE = re.compile(r"(\w)-\n(\w)")


def _pdf_paragraphs(page_text: str) -> list[str]:
    """Rebuild paragraphs from pypdf's line-oriented page text."""
    text = page_text.replace("\r\n", "\n").replace("\r", "\n")
    text = _DEHYPHEN_RE.sub(
        lambda m: m.group(1) + m.group(2) if m.group(2).islower() else m.group(0), text
    )
    lines = text.split("\n")
    widths = [len(line.strip()) for line in lines if line.strip()]
    typical = sorted(widths)[int(len(widths) * 0.8)] if widths else 0
    paragraphs: list[str] = []
    current: list[str] = []
    for raw in lines:
        line = raw.strip()
        if not line:
            if current:
                paragraphs.append(" ".join(current))
                current = []
            continue
        current.append(line)
        ends_sentence = line[-1] in '.!?:"”’)'
        short = typical > 0 and len(line) < typical * 0.7
        if ends_sentence and short:
            paragraphs.append(" ".join(current))
            current = []
    if current:
        paragraphs.append(" ".join(current))
    return paragraphs


def extract_pdf(data: bytes) -> Extracted:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:
                raise TamberError(
                    422, "extract_failed", "The PDF is encrypted and can't be read."
                ) from exc
        paragraphs: list[str] = []
        for page in reader.pages:
            try:
                paragraphs.extend(_pdf_paragraphs(page.extract_text() or ""))
            except Exception:
                logger.debug("pypdf could not read a page", exc_info=True)
        title: str | None = None
        language: str | None = None
        try:
            meta = reader.metadata
            if meta is not None and meta.title:
                title = clean_paragraph(str(meta.title)) or None
        except Exception:
            title = None
        try:
            root: Any = reader.trailer["/Root"]
            lang = root.get("/Lang")
            if lang:
                language = str(lang).strip() or None
        except Exception:
            language = None
    except TamberError:
        raise
    except (PdfReadError, ValueError, KeyError, OSError) as exc:
        raise TamberError(422, "extract_failed", "The PDF could not be read.") from exc
    return Extracted(title=title, paragraphs=paragraphs, language=language, mime_type=PDF)


def extract_docx(data: bytes) -> Extracted:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise TamberError(422, "extract_failed", "The DOCX file could not be read.") from exc
    paragraphs: list[str] = []
    heading: str | None = None

    def add_paragraph(p: Any) -> None:
        nonlocal heading
        text = p.text or ""
        if not text.strip():
            return
        style = (getattr(getattr(p, "style", None), "name", "") or "").lower()
        if heading is None and (style.startswith("heading") or style == "title"):
            heading = clean_paragraph(text) or None
        paragraphs.append(text)

    def add_table(table: Any) -> None:
        for row in table.rows:
            seen: set[int] = set()
            for cell in row.cells:
                if id(cell._tc) in seen:
                    continue  # merged cells repeat
                seen.add(id(cell._tc))
                for inner in cell.iter_inner_content():
                    if isinstance(inner, Paragraph):
                        add_paragraph(inner)
                    elif isinstance(inner, Table):
                        add_table(inner)

    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            add_paragraph(block)
        elif isinstance(block, Table):
            add_table(block)
    props = document.core_properties
    title = clean_paragraph(props.title or "") or heading
    language = (props.language or "").strip() or None
    return Extracted(title=title or None, paragraphs=paragraphs, language=language, mime_type=DOCX)


def extract_epub(data: bytes) -> Extracted:
    import ebooklib
    from ebooklib import epub

    try:
        book = epub.read_epub(io.BytesIO(data), options={"ignore_ncx": True})
    except Exception as exc:
        raise TamberError(422, "extract_failed", "The EPUB file could not be read.") from exc
    paragraphs: list[str] = []
    heading: str | None = None
    for idref, _linear in book.spine:
        item = book.get_item_with_id(idref)
        if item is None or item.get_type() != ebooklib.ITEM_DOCUMENT:
            continue
        soup = BeautifulSoup(item.get_content(), "lxml")
        if heading is None:
            heading = first_heading(soup)
        paragraphs.extend(block_paragraphs(soup.body or soup))
    title_meta = book.get_metadata("DC", "title")
    title = clean_paragraph(str(title_meta[0][0])) if title_meta else None
    lang_meta = book.get_metadata("DC", "language")
    language = str(lang_meta[0][0]).strip() if lang_meta else None
    return Extracted(
        title=title or heading, paragraphs=paragraphs, language=language or None, mime_type=EPUB
    )


def extract_markdown(data: bytes) -> Extracted:
    from markdown_it import MarkdownIt

    source = decode_text(data)
    md = MarkdownIt("commonmark").enable("table").enable("strikethrough")
    rendered = md.render(source)
    soup = BeautifulSoup(rendered, "lxml")
    heading = first_heading(soup)
    paragraphs = block_paragraphs(soup.body or soup, drop_code=True)
    return Extracted(title=heading, paragraphs=paragraphs, language=None, mime_type=MARKDOWN)


def extract_text_file(data: bytes) -> Extracted:
    text = decode_text(data)
    return Extracted(
        title=None, paragraphs=split_plain_paragraphs(text), language=None, mime_type=TEXT
    )


def extract_html_file(data: bytes, url: str | None = None) -> Extracted:
    doc = extract_html(data, url)
    return Extracted(
        title=doc.title or doc.first_heading,
        paragraphs=doc.paragraphs,
        language=doc.language,
        mime_type=HTML,
    )


def extract_document(
    data: bytes, content_type: str | None, filename: str | None, url: str | None = None
) -> Extracted:
    kind = detect_type(data, content_type, filename)
    if kind == PDF:
        return extract_pdf(data)
    if kind == DOCX:
        return extract_docx(data)
    if kind == EPUB:
        return extract_epub(data)
    if kind == HTML:
        return extract_html_file(data, url)
    if kind == MARKDOWN:
        return extract_markdown(data)
    if kind == TEXT:
        return extract_text_file(data)
    raise TamberError(
        415,
        "unsupported_media_type",
        "Unsupported document type. Supported: PDF, DOCX, EPUB, HTML, Markdown and plain text.",
        param="file",
    )
