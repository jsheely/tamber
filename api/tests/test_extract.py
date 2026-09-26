"""POST /v1/extract: files, HTML, URLs (mocked transport), SSRF guard, normalisation."""

from __future__ import annotations

import io
from pathlib import Path
from typing import Any

import httpx
import pytest
from fastapi.testclient import TestClient

from tamber_api.extract.normalize import finalize, split_plain_paragraphs
from tamber_api.extract.ssrf import ip_is_blocked

from .conftest import make_app, ready_client, wait_ready

ARTICLE_HTML = """<!doctype html>
<html lang="en-GB"><head><title>Why Voices Matter | Example</title>
<meta property="og:title" content="Why Voices Matter"></head>
<body>
<nav><a href="/">Home</a> <a href="/about">About</a></nav>
<article>
<h1>Why Voices Matter</h1>
<p>The first paragraph of the article, with   extra   spaces and a soft­hyphen.</p>
<p>The second paragraph talks about how listening to long-form writing changes the way we
remember it, and why a natural voice makes that experience far more pleasant.</p>
<ul><li>First point worth hearing.</li><li>Second point worth hearing.</li></ul>
<p>A closing paragraph with enough words to count as real article content for the extractor.</p>
</article>
<footer>Copyright notice</footer>
</body></html>"""


def check_output_rules(body: dict[str, Any]) -> None:
    text = body["text"]
    assert text == text.strip()
    assert "­" not in text and "​" not in text
    for para in text.split("\n\n"):
        assert para and para == para.strip()
        assert "  " not in para and "\n" not in para and "\t" not in para
    assert body["word_count"] == len(text.split())
    assert body["char_count"] == len(text.encode("utf-16-le")) // 2


def minimal_pdf(lines: list[str], title: str | None = None) -> bytes:
    """A tiny valid PDF with one page of Helvetica text (built by hand, no dependency)."""
    content_ops = ["BT", "/F1 12 Tf", "14 TL", "72 720 Td"]
    for line in lines:
        escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        content_ops.append(f"({escaped}) Tj T*")
    content_ops.append("ET")
    stream = "\n".join(content_ops).encode("latin-1")
    info = f"<< /Title ({title}) >>" if title else "<< >>"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R /Lang (en-US) >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        info.encode("latin-1"),
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R /Info 6 0 R >>\n"
        f"startxref\n{xref}\n%%EOF\n".encode()
    )
    return out.getvalue()


def make_docx() -> bytes:
    import docx

    document = docx.Document()
    document.core_properties.language = "en-US"
    document.add_heading("Quarterly Notes", level=1)
    document.add_paragraph("The first   paragraph of the document.")
    document.add_paragraph("")
    document.add_paragraph("Second paragraph, still going.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Cell one text"
    table.rows[0].cells[1].text = "Cell two text"
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def make_epub(tmp_path: Path) -> bytes:
    from ebooklib import epub

    book = epub.EpubBook()
    book.set_identifier("tamber-test")
    book.set_title("A Tiny Book")
    book.set_language("en")
    ch2 = epub.EpubHtml(title="Two", file_name="two.xhtml", lang="en")
    ch2.content = "<h1>Chapter Two</h1><p>Second chapter text.</p>"
    ch1 = epub.EpubHtml(title="One", file_name="one.xhtml", lang="en")
    ch1.content = "<h1>Chapter One</h1><p>First chapter text.</p>"
    book.add_item(ch2)
    book.add_item(ch1)
    book.toc = [ch1, ch2]
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    book.spine = [ch1, ch2]  # spine order differs from insertion order
    path = tmp_path / "book.epub"
    epub.write_epub(str(path), book)
    return path.read_bytes()


def upload(client: TestClient, name: str, data: bytes, ctype: str, **extra: str) -> Any:
    return client.post("/v1/extract", files={"file": (name, data, ctype)}, data=extra)


def test_txt_upload(client: TestClient) -> None:
    data = "Title line\r\nwraps here.\r\n\r\n\r\nSecond\tparagraph​.\n".encode()
    r = upload(client, "notes.txt", data, "text/plain")
    assert r.status_code == 200, r.text
    body = r.json()
    check_output_rules(body)
    assert body["text"] == "Title line wraps here.\n\nSecond paragraph."
    assert body["title"] == "notes"
    assert body["source"] == "notes.txt" and body["source_type"] == "file"
    assert body["mime_type"] == "text/plain" and body["truncated"] is False


def test_txt_non_utf8_is_detected(client: TestClient) -> None:
    data = "Café au lait, s'il vous plaît. Très bien, merci beaucoup.".encode("cp1252")
    r = upload(client, "fr.txt", data, "application/octet-stream")
    assert r.status_code == 200
    assert "Café" in r.json()["text"]


def test_markdown_upload_drops_code_blocks(client: TestClient) -> None:
    md = (
        "# Read Me\n\nSome *emphasis* and a [link text](https://x.example).\n\n"
        "```python\nprint('code is not read')\n```\n\n- item one\n- item two\n\n"
        "Inline `code` stays.\n"
    )
    r = upload(client, "README.md", md.encode(), "application/octet-stream")
    body = r.json()
    check_output_rules(body)
    assert body["title"] == "Read Me"
    assert body["text"] == (
        "Read Me\n\nSome emphasis and a link text.\n\nitem one\n\nitem two\n\nInline code stays."
    )
    assert "print" not in body["text"]
    assert body["mime_type"] == "text/markdown"


def test_html_upload_and_html_json(client: TestClient) -> None:
    r = upload(client, "page.html", ARTICLE_HTML.encode(), "text/html")
    body = r.json()
    check_output_rules(body)
    assert body["title"] == "Why Voices Matter"
    assert "first paragraph of the article, with extra spaces and a softhyphen." in body["text"]
    assert "Home" not in body["text"] and "Copyright" not in body["text"]
    assert body["language"] == "en-GB"
    r = client.post("/v1/extract", json={"html": ARTICLE_HTML, "url": "https://example.com/a"})
    body = r.json()
    check_output_rules(body)
    assert body["source"] == "https://example.com/a" and body["source_type"] == "html"
    assert body["mime_type"] == "text/html"
    assert "Second point worth hearing." in body["text"].split("\n\n")
    r = client.post("/v1/extract", json={"html": ARTICLE_HTML})
    assert r.json()["source"] == ""


def test_html_fallback_when_trafilatura_finds_nothing(client: TestClient) -> None:
    html = "<html><body><div>Just one short line.</div></body></html>"
    r = client.post("/v1/extract", json={"html": html})
    assert r.status_code == 200
    assert r.json()["text"] == "Just one short line."


def test_docx_upload(client: TestClient) -> None:
    r = upload(client, "notes.docx", make_docx(), "application/octet-stream")
    assert r.status_code == 200, r.text
    body = r.json()
    check_output_rules(body)
    assert body["title"] == "Quarterly Notes"
    assert body["text"].split("\n\n") == [
        "Quarterly Notes",
        "The first paragraph of the document.",
        "Second paragraph, still going.",
        "Cell one text",
        "Cell two text",
    ]
    assert body["language"] == "en-US"


def test_docx_filename_override(client: TestClient) -> None:
    r = upload(client, "blob", make_docx(), "application/octet-stream", filename="Report.docx")
    assert r.status_code == 200
    assert r.json()["source"] == "Report.docx"


def test_pdf_upload(client: TestClient) -> None:
    pdf = minimal_pdf(
        ["Hello from a tiny PDF document.", "It has a second line of text."], title="Tiny PDF"
    )
    r = upload(client, "tiny.pdf", pdf, "application/pdf")
    assert r.status_code == 200, r.text
    body = r.json()
    check_output_rules(body)
    assert body["title"] == "Tiny PDF"
    assert "Hello from a tiny PDF document." in body["text"]
    assert "second line of text." in body["text"]
    assert body["mime_type"] == "application/pdf"
    assert body["language"] == "en-US"
    # Detected by magic bytes without a helpful name or type.
    r = upload(client, "download", pdf, "application/octet-stream")
    assert r.status_code == 200 and r.json()["mime_type"] == "application/pdf"


def test_pdf_without_text_is_422(client: TestClient) -> None:
    r = upload(client, "scan.pdf", minimal_pdf([]), "application/pdf")
    assert r.status_code == 422
    assert r.json()["error"]["code"] == "extract_failed"


def test_epub_spine_order(client: TestClient, tmp_path: Path) -> None:
    r = upload(client, "book.epub", make_epub(tmp_path), "application/epub+zip")
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["title"] == "A Tiny Book"
    text = body["text"]
    assert text.index("Chapter One") < text.index("Chapter Two")
    assert body["language"] == "en"


def test_unsupported_and_invalid_requests(client: TestClient) -> None:
    r = upload(client, "image.png", b"\x89PNG\r\n\x1a\n\x00\x00\x00", "image/png")
    assert r.status_code == 415 and r.json()["error"]["code"] == "unsupported_media_type"
    r = client.post("/v1/extract", json={})
    assert r.status_code == 400 and r.json()["error"]["code"] == "invalid_request"
    r = client.post("/v1/extract", files={"other": ("a.txt", b"hi", "text/plain")})
    assert r.status_code == 400
    r = client.post("/v1/extract", json={"html": "<html><body></body></html>"})
    assert r.status_code == 422 and r.json()["error"]["code"] == "extract_failed"
    r = client.post("/v1/extract", content=b"hi", headers={"Content-Type": "image/png"})
    assert r.status_code == 415


@pytest.mark.parametrize(
    "url",
    [
        "http://127.0.0.1/",
        "http://10.0.0.1/admin",
        "http://[::1]:8080/",
        "http://169.254.169.254/latest/meta-data/",
        "http://100.64.1.2/",  # CGNAT (NetBird peers)
        "http://0.0.0.0/",
        "http://localhost:8880/v1/health",
        "ftp://example.com/file.txt",
        "file:///etc/passwd",
        "http://user:pass@example.com/",
    ],
)
def test_ssrf_rejections(client: TestClient, url: str) -> None:
    r = client.post("/v1/extract", json={"url": url})
    assert r.status_code == 400, r.text
    assert r.json()["error"]["code"] == "url_not_allowed"
    assert r.json()["error"]["param"] == "url"


def test_ip_classification() -> None:
    import ipaddress

    assert ip_is_blocked(ipaddress.ip_address("::ffff:127.0.0.1"))
    assert ip_is_blocked(ipaddress.ip_address("192.168.1.10"))
    assert not ip_is_blocked(ipaddress.ip_address("93.184.216.34"))
    assert not ip_is_blocked(ipaddress.ip_address("2606:4700:4700::1111"))


def _url_client(handler: Any, **overrides: Any) -> Any:
    app = make_app(**overrides)
    app.state.extract_transport = httpx.MockTransport(handler)
    return app


def test_url_fetch_redirects_and_dispatch() -> None:
    pdf = minimal_pdf(["A PDF served from a URL."])
    seen: list[str] = []

    def handler(request: httpx.Request) -> httpx.Response:
        seen.append(str(request.url))
        assert request.headers["user-agent"].startswith("Tamber/")
        if request.url.path == "/start":
            return httpx.Response(302, headers={"Location": "/paper"})
        if request.url.path == "/paper":
            return httpx.Response(200, content=pdf, headers={"Content-Type": "application/pdf"})
        if request.url.path == "/article":
            return httpx.Response(
                200, content=ARTICLE_HTML.encode(), headers={"Content-Type": "text/html"}
            )
        return httpx.Response(404)

    app = _url_client(handler, extract_allow_private=True)
    with TestClient(app) as client:
        wait_ready(client)
        r = client.post("/v1/extract", json={"url": "http://docs.test/start"})
        assert r.status_code == 200, r.text
        body = r.json()
        assert body["text"] == "A PDF served from a URL."
        assert body["source"] == "http://docs.test/start" and body["source_type"] == "url"
        assert body["mime_type"] == "application/pdf" and body["title"] == "paper"
        assert seen == ["http://docs.test/start", "http://docs.test/paper"]
        r = client.post("/v1/extract", json={"url": "http://docs.test/article"})
        assert r.json()["title"] == "Why Voices Matter"
        r = client.post("/v1/extract", json={"url": "http://docs.test/missing"})
        assert r.status_code == 502
        assert r.json()["error"]["code"] == "fetch_failed"
        assert r.json()["error"]["details"] == {"status": 404}


def test_url_redirect_into_private_range_is_blocked() -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(302, headers={"Location": "http://127.0.0.1:8880/v1/health"})

    app = _url_client(handler)
    with TestClient(app) as client:
        wait_ready(client)
        r = client.post("/v1/extract", json={"url": "http://93.184.216.34/"})
        assert r.status_code == 400
        assert r.json()["error"]["code"] == "url_not_allowed"


def test_url_too_many_redirects_and_size_cap() -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        if request.url.path.startswith("/loop"):
            return httpx.Response(302, headers={"Location": "/loop" + request.url.path[5:] + "x"})
        return httpx.Response(200, content=b"a " * 400_000, headers={"Content-Type": "text/plain"})

    app = _url_client(handler, extract_allow_private=True, extract_max_download_mb=0.5)
    with TestClient(app) as client:
        wait_ready(client)
        r = client.post("/v1/extract", json={"url": "http://docs.test/loop"})
        assert r.status_code == 502 and "redirects" in r.json()["error"]["message"]
        r = client.post("/v1/extract", json={"url": "http://docs.test/big.txt"})
        assert r.status_code == 413 and r.json()["error"]["code"] == "file_too_large"


def test_truncation_at_boundaries() -> None:
    paragraphs = [f"Paragraph {i} has a few words in it." for i in range(50)]
    out = finalize(paragraphs, 100)
    assert out.truncated and out.char_count <= 100
    assert out.text.endswith("in it.") and "\n\n" in out.text
    one = finalize(["First sentence here. Second sentence is longer than the rest of it."], 40)
    assert one.truncated and one.text == "First sentence here."
    words = finalize(["wordy " * 40], 50)
    assert words.truncated and not words.text.endswith(" ") and words.char_count <= 50
    with ready_client(max_extract_chars=100) as client:
        r = upload(client, "long.txt", "\n\n".join(paragraphs).encode(), "text/plain")
        body = r.json()
        assert body["truncated"] is True and body["char_count"] <= 100


def test_plain_paragraph_split() -> None:
    assert split_plain_paragraphs("a\nb\n\n \n c\r\n\r\nd") == ["a\nb", "c", "d"]
